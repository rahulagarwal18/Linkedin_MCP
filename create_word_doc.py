import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

doc = Document()

# Page Setup: Normal 1-inch margins
for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

# Color Palette: Executive Blue theme
PRIMARY = RGBColor(10, 102, 194)    # LinkedIn / Executive Blue #0A66C2
SECONDARY = RGBColor(27, 42, 74)   # Deep Navy #1B2A4A
DARK_TEXT = RGBColor(33, 37, 41)   # #212529
MUTED_GRAY = RGBColor(108, 117, 125)

def format_run(run, font_name="Calibri", font_size=11, color=DARK_TEXT, bold=False, italic=False):
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.color.rgb = color
    run.bold = bold
    run.italic = italic

# Title
title_p = doc.add_paragraph()
title_p.paragraph_format.space_before = Pt(0)
title_p.paragraph_format.space_after = Pt(4)
run_title = title_p.add_run("Connecting AI with LinkedIn Safely")
format_run(run_title, font_name="Segoe UI", font_size=24, color=PRIMARY, bold=True)

# Subtitle
sub_p = doc.add_paragraph()
sub_p.paragraph_format.space_after = Pt(16)
run_sub = sub_p.add_run("A Guardrail Architecture for Claude & Antigravity (Without Account Restrictions)")
format_run(run_sub, font_name="Segoe UI", font_size=13, color=MUTED_GRAY, italic=True)

# Infographic Image
image_path = r"C:\Users\rahul\.gemini\antigravity\brain\41d62ef6-c4e2-4337-b09a-dccece622e7f\linkedin_mcp_architecture_1791780623719.jpg"
if os.path.exists(image_path):
    img_p = doc.add_paragraph()
    img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    img_p.paragraph_format.space_before = Pt(8)
    img_p.paragraph_format.space_after = Pt(4)
    run_img = img_p.add_run()
    run_img.add_picture(image_path, width=Inches(6.2))

    cap_p = doc.add_paragraph()
    cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap_p.paragraph_format.space_after = Pt(20)
    run_cap = cap_p.add_run("Figure 1: How the Safe Guardrail MCP Server protects accounts with human pacing and persistent profiles.")
    format_run(run_cap, font_name="Segoe UI", font_size=9.5, color=MUTED_GRAY, italic=True)

def add_heading_1(text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(16)
    h.paragraph_format.space_after = Pt(6)
    r = h.add_run(text)
    format_run(r, font_name="Segoe UI", font_size=15, color=SECONDARY, bold=True)
    return h

def add_body(text, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text)
    format_run(r, font_name="Calibri", font_size=11, color=DARK_TEXT)
    return p

def add_bullet(lead, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(4)
    r_lead = p.add_run(lead + " ")
    format_run(r_lead, font_name="Calibri", font_size=11, color=DARK_TEXT, bold=True)
    r_text = p.add_run(text)
    format_run(r_text, font_name="Calibri", font_size=11, color=DARK_TEXT)
    return p

# Section 1
add_heading_1("1. The Dilemma with Existing LinkedIn Tools")
add_body(
    "When professionals connect an AI assistant like Claude or Antigravity to LinkedIn, they immediately hit a wall: "
    "there is no official LinkedIn connector available. As a result, users are forced to choose between two flawed alternatives:"
)
add_bullet("The Official API Tools:", "Safe from bans, but very limited. You can publish a basic text post, but you cannot search jobs, read user profiles, or send messages.")
add_bullet("The Raw Scraper Tools:", "Can do everything, but act like high-speed bots. They demand your private session cookies, blast requests all at once, and quickly trigger LinkedIn's automated security flags and account bans.")

add_body(
    "We built this custom connector so nobody has to compromise between 'safe but useless' and 'powerful but banned'."
)

# Section 2
add_heading_1("2. What We Actually Built")
add_body(
    "We designed a custom connector (Model Context Protocol / MCP server) that runs locally on your machine. "
    "Instead of blasting LinkedIn with rapid-fire commands, it behaves like a careful human assistant sitting in front of a real browser:"
)
add_bullet("Persistent Browser Profile:", "Uses a standard local Chrome profile that stays logged in naturally, avoiding suspicious re-logins.")
add_bullet("Human-Paced Action Queue:", "Any outbound networking request (message, connection request) enters a waiting queue with randomized pauses (35 to 95 seconds apart).")
add_bullet("Hard Daily Safety Limits:", "Monitors your daily activity and stops automatically before reaching spam thresholds.")

# Section 3
add_heading_1("3. Why This Is Better & Safer Than Typical 3rd-Party Tools")
add_bullet("No Leaked or Pasted Cookies:", "Typical third-party tools require you to open developer tools, extract secret cookies (li_at), and paste them into configuration files. Our solution opens a clean browser window once; you log in naturally with your password and 2FA, and the session stays securely on your machine.")
add_bullet("Human Rhythm (Anti-Burst Protection):", "When standard automation sends 5 connection requests, it fires all 5 within 2 seconds. Our queue spaces each request with natural, randomized human delays so it never looks like robotic spam.")
add_bullet("Built-in Speed Governors:", "If an AI accidentally loops or repeats a command, our server has strict daily ceilings (maximum 15 invites, 15 messages, and 50 profile views per day). Once reached, it politely pauses until the next day.")
add_bullet("Universal AI Compatibility:", "Works out-of-the-box with Claude Desktop, Antigravity, Cursor, or any platform supporting the MCP standard.")

# Section 4
add_heading_1("4. Core Features Included")
add_bullet("Job Search:", "Search job listings by title, keywords, and location (e.g., 'Senior AI Engineer in Remote').")
add_bullet("Profile Reader:", "Inspect public profile headlines, summaries, and career backgrounds before meetings.")
add_bullet("Post Drafting & Scheduling:", "Draft and schedule posts directly to your LinkedIn feed without manual copying and pasting.")
add_bullet("Paced Direct Messaging:", "Queue messages to contacts that send smoothly at human intervals.")
add_bullet("Connection Requests with Notes:", "Send personalized invitation notes paced safely one by one.")
add_bullet("Safety Dashboard:", "Ask the AI at any time to report how much daily quota has been used and what remains.")

# Section 5
add_heading_1("5. The Honest Truth About Safety (No False Claims)")
add_body(
    "We believe in complete transparency, so here is what you should honestly know:"
)
add_bullet("Not Officially Endorsed:", "This is not an official LinkedIn product. LinkedIn does not officially support automated browser tools on personal accounts.")
add_bullet("No 100% Zero-Risk Guarantee:", "No automation software can claim absolute zero risk, because LinkedIn continually updates its anti-bot algorithms.")
add_bullet("The Safest Practical Choice:", "By keeping your login local, using realistic human delays, and enforcing strict daily quotas, this architecture reduces risk to the absolute minimum feasible level.")

# Save document
output_path = r"B:\AI_Projects\linkedin-mcp-agent\Connecting_AI_with_LinkedIn_Safely.docx"
doc.save(output_path)
print(f"Document saved successfully to {output_path}")
